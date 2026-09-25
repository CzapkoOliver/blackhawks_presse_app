"""
nachbericht_docx.py
---------------------
Wandelt den von nachbericht_generator.py erzeugten Markdown-Text in ein
Word-Dokument (.docx) um.

Liegt die Datei "vorlage_nachbericht.docx" im selben Ordner (die vom Verein
bereitgestellte Layout-Vorlage mit Briefkopf, Logos und Fusszeile), wird
GENAU DIESE Datei als Basis verwendet: der bisherige Inhalt wird entfernt,
Kopf-/Fusszeile, Logos, Seitenraender und Schriftart der Vorlage bleiben
dabei vollstaendig erhalten. So muss der fertige Nachbericht nicht mehr per
Hand in die Vorlage kopiert werden. Fehlt die Vorlagen-Datei (z.B. beim
allerersten Test ohne Upload), wird ersatzweise ein einfaches, neutrales
Dokument erzeugt.

Bewusst ein einfacher, zeilenbasierter Markdown-Parser statt einer
allgemeinen Markdown-Bibliothek - der Generator erzeugt immer dieselbe,
bekannte Struktur, dafuer reicht das aus und bleibt gut wartbar.
"""

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# Vereinsfarbe Rot (Schwarz/Rot/Weiss), passend zur Vorlage.
VEREINSROT = RGBColor(0xC8, 0x10, 0x2E)

VORLAGE_PFAD = Path(__file__).resolve().parent / "vorlage_nachbericht.docx"


def _fuege_inline_formatierten_text_hinzu(absatz, text: str, fett_standard: bool = False) -> None:
    """
    Fuegt Text zu einem Absatz hinzu und beachtet dabei **fett**-Markierungen
    innerhalb der Zeile (z.B. "**Name:** Zitat").
    """
    teile = re.split(r"(\*\*.*?\*\*)", text)
    for teil in teile:
        if not teil:
            continue
        if teil.startswith("**") and teil.endswith("**"):
            lauf = absatz.add_run(teil[2:-2])
            lauf.bold = True
        else:
            lauf = absatz.add_run(teil)
            lauf.bold = fett_standard


def _leere_dokument_body(dokument: Document) -> None:
    """
    Entfernt alle vorhandenen Absaetze aus dem Dokument-Body, laesst dabei
    aber die abschliessenden Seiteneinstellungen (sectPr) unangetastet -
    genau die enthalten den Verweis auf Kopf-/Fusszeile, Logos und
    Seitenraender der Vorlage.
    """
    body = dokument.element.body
    for kind in list(body):
        if kind.tag.endswith("}sectPr"):
            continue
        body.remove(kind)


def _neues_dokument() -> Document:
    if VORLAGE_PFAD.exists():
        dokument = Document(str(VORLAGE_PFAD))
        _leere_dokument_body(dokument)
        return dokument

    # Fallback ohne Vorlage (z.B. falls vorlage_nachbericht.docx noch nicht
    # mit hochgeladen wurde) - einfaches, neutrales Dokument statt Absturz.
    dokument = Document()
    stil = dokument.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(11)
    return dokument


def nachbericht_zu_docx_bytes(nachbericht_text: str) -> bytes:
    dokument = _neues_dokument()

    zeilen = nachbericht_text.strip().splitlines()

    for rohe_zeile in zeilen:
        zeile = rohe_zeile.strip()

        if not zeile:
            continue

        if zeile.lower().startswith("sperrfrist bis"):
            # Wie in der Vorlage: "Sperrfrist bis:" auf einer eigenen,
            # unterstrichenen Zeile, der Wert (meist "keine") darunter in
            # der Vereinsfarbe Rot.
            label, _, wert = zeile.partition(":")
            absatz_label = dokument.add_paragraph()
            lauf_label = absatz_label.add_run(f"{label.strip()}:")
            lauf_label.bold = True
            lauf_label.underline = True

            absatz_wert = dokument.add_paragraph()
            lauf_wert = absatz_wert.add_run(wert.strip() or "keine")
            lauf_wert.bold = True
            lauf_wert.font.color.rgb = VEREINSROT
            continue

        if zeile.startswith("# "):
            absatz = dokument.add_paragraph()
            absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lauf = absatz.add_run(zeile[2:].strip())
            lauf.bold = True
            lauf.font.size = Pt(16)
            continue

        if zeile.startswith("## "):
            absatz = dokument.add_paragraph()
            absatz.paragraph_format.space_before = Pt(12)
            lauf = absatz.add_run(zeile[3:].strip())
            lauf.underline = True
            lauf.font.size = Pt(12)
            continue

        # Eigene Zeile, die komplett fett ist (z.B. die reisserische Vorzeile
        # vor der Headline): "**Vorzeile**"
        if zeile.startswith("**") and zeile.endswith("**") and zeile.count("**") == 2:
            absatz = dokument.add_paragraph()
            absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lauf = absatz.add_run(zeile[2:-2])
            lauf.bold = True
            continue

        absatz = dokument.add_paragraph()
        _fuege_inline_formatierten_text_hinzu(absatz, zeile)

    puffer = BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()
